"""
Generates the multi-sheet Excel evaluation report:
  1. Results               - one row per prediction (incl. bundle children)
  2. Error Analysis         - only incorrect predictions
  3. Per-Class Accuracy      - accuracy grouped by ground-truth class
  4. RAG vs Gemini           - agreement/override analysis
  5. RAG vs Gemini Dependency - DOCUMENT-LEVEL summary: how many physical
                               files were resolved by RAG alone vs needed
                               Gemini for their final answer, counted once
                               per file (not once per output row) — a
                               multi-page bundle document counts as ONE
                               document here, unlike the Results sheet
                               above where it's split into child rows.
  6. Configuration           - run settings + cost/quota summary
"""
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def _compute_document_level_dependency(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    """
    Groups output rows by their ORIGINAL FILE PATH so each physical document
    is counted exactly once, regardless of how many child rows a bundle
    document produced.

    A document counts as "Gemini-dependent" if EITHER:
      - It's a bundle document (Layer == 2 for any of its rows) — bundles
        always rely on Gemini for the final child-level answer, even if
        Layer 1 identified the bundle itself via RAG.
      - It's a standalone document whose Method involved gemini
        (gemini / gemini_only / gemini_forced).
    Otherwise it counts as "RAG-only".
    """
    by_doc: Dict[str, List[Dict[str, Any]]] = {}
    for row in rows:
        path = row.get("Document_Path", "UNKNOWN_PATH")
        by_doc.setdefault(path, []).append(row)

    rag_only_docs = 0
    gemini_docs = 0

    for path, doc_rows in by_doc.items():
        is_bundle = any(r.get("Layer") == 2 for r in doc_rows)
        if is_bundle:
            gemini_docs += 1
            continue

        # Standalone: single row, check its Method
        method = str(doc_rows[0].get("Method", ""))
        if "gemini" in method:
            gemini_docs += 1
        else:
            rag_only_docs += 1

    total_docs = rag_only_docs + gemini_docs
    rag_pct = round(rag_only_docs / total_docs * 100, 1) if total_docs else 0.0
    gemini_pct = round(gemini_docs / total_docs * 100, 1) if total_docs else 0.0

    return {
        "total_documents": total_docs,
        "rag_only_documents": rag_only_docs,
        "gemini_dependent_documents": gemini_docs,
        "rag_only_pct": rag_pct,
        "gemini_dependent_pct": gemini_pct,
    }


def _style_header(ws, freeze_col="A2"):
    ws.freeze_panes = freeze_col
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for col_idx, cell in enumerate(ws[1], 1):
        col_letter = get_column_letter(col_idx)
        header_len = len(str(cell.value or ""))
        ws.column_dimensions[col_letter].width = min(max(header_len + 4, 12), 40)

    if ws.title == "Results":
        red_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
        green_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
        is_correct_col = None
        for col_idx, cell in enumerate(ws[1], 1):
            if cell.value == "Is_Correct":
                is_correct_col = col_idx
                break
        if is_correct_col:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                val = row[is_correct_col - 1].value
                fill = green_fill if val is True else red_fill if val is False else None
                if fill:
                    for cell in row:
                        cell.fill = fill

    if ws.title == "Error Analysis":
        red_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.fill = red_fill


def build_report(
    rows: List[Dict[str, Any]],
    output_dir: str,
    run_config_summary: Dict[str, Any],
    gemini_usage_summary: Dict[str, Any],
    dataset_folder: str,
) -> str:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_file = output_dir / f"classification_report_{timestamp}.xlsx"

    df = pd.DataFrame(rows)

    total_docs = len(df)
    correct_count = int(df["Is_Correct"].sum()) if total_docs else 0
    error_count = int((df["Predicted_Class"] == "ERROR").sum()) if total_docs else 0
    success_count = total_docs - error_count
    overall_accuracy = round((correct_count / total_docs) * 100, 1) if total_docs else 0.0

    with pd.ExcelWriter(excel_file, engine="openpyxl") as writer:
        # ---------------- Sheet 1: Results ----------------
        if total_docs:
            df.to_excel(writer, index=False, sheet_name="Results")
        else:
            pd.DataFrame([{"Message": "No documents were classified."}]).to_excel(
                writer, index=False, sheet_name="Results"
            )
        _style_header(writer.sheets["Results"])

        # ---------------- Sheet 2: Error Analysis ----------------
        if total_docs:
            error_df = df[df["Is_Correct"] == False][  # noqa: E712
                [c for c in [
                    "Document_Name", "Actual_Class", "Predicted_Class",
                    "RAG_Class", "RAG_Confidence", "Gemini_Class", "Gemini_Confidence",
                    "Method", "Error",
                ] if c in df.columns]
            ]
            if len(error_df):
                error_df.to_excel(writer, index=False, sheet_name="Error Analysis")
            else:
                pd.DataFrame([{"Message": "No errors — 100% accuracy on this run."}]).to_excel(
                    writer, index=False, sheet_name="Error Analysis"
                )
            _style_header(writer.sheets["Error Analysis"])

        # ---------------- Sheet 3: Per-Class Accuracy ----------------
        if total_docs:
            stats = df.groupby("Actual_Class").agg({"Is_Correct": ["sum", "count"]}).reset_index()
            stats.columns = ["Class", "Correct", "Total"]
            stats["Accuracy_%"] = (stats["Correct"] / stats["Total"] * 100).round(1)
            stats = stats.sort_values("Accuracy_%")
            stats.to_excel(writer, index=False, sheet_name="Per-Class Accuracy")
            _style_header(writer.sheets["Per-Class Accuracy"])

        # ---------------- Sheet 4: RAG vs Gemini ----------------
        if total_docs and "RAG_Class" in df.columns and "Gemini_Class" in df.columns:
            agree_df = df[(df["RAG_Class"] != "N/A") & (df["Gemini_Class"] != "N/A")].copy()
            if len(agree_df):
                agree_df["Agree"] = agree_df["RAG_Class"] == agree_df["Gemini_Class"]
                summary_rows = [
                    {"Metric": "Total Docs with Both RAG & Gemini", "Value": len(agree_df)},
                    {"Metric": "RAG & Gemini Agree", "Value": int(agree_df["Agree"].sum())},
                    {"Metric": "RAG & Gemini Disagree", "Value": int((~agree_df["Agree"]).sum())},
                    {"Metric": "Agreement Rate %", "Value": round(agree_df["Agree"].mean() * 100, 1)},
                ]
                pd.DataFrame(summary_rows).to_excel(writer, index=False, sheet_name="RAG vs Gemini", startrow=0)

                disagree_df = agree_df[~agree_df["Agree"]][
                    [c for c in [
                        "Document_Name", "Actual_Class", "RAG_Class", "RAG_Confidence",
                        "Gemini_Class", "Gemini_Confidence", "Predicted_Class", "Is_Correct",
                    ] if c in agree_df.columns]
                ]
                if len(disagree_df):
                    disagree_df.to_excel(
                        writer, index=False, sheet_name="RAG vs Gemini", startrow=len(summary_rows) + 3
                    )
                _style_header(writer.sheets["RAG vs Gemini"])

        # ---------------- Sheet 5: RAG vs Gemini Dependency (document-level) ----------------
        dep = _compute_document_level_dependency(rows)
        dependency_rows = [
            {"Metric": "Total Physical Documents", "Value": dep["total_documents"]},
            {"Metric": "Resolved by RAG Alone", "Value": dep["rag_only_documents"]},
            {"Metric": "Needed Gemini (incl. all bundle documents)", "Value": dep["gemini_dependent_documents"]},
            {"Metric": "RAG-Only %", "Value": dep["rag_only_pct"]},
            {"Metric": "Gemini-Dependent %", "Value": dep["gemini_dependent_pct"]},
            {"Metric": "", "Value": ""},
            {"Metric": "Note", "Value": (
                "Counted per physical FILE, not per output row. Bundle documents "
                "(e.g. MOA/AOA/COI files) always count as Gemini-dependent here, "
                "since their final child-level answer always comes from Layer 2's "
                "Gemini-only page splitting, even if Layer 1 identified the bundle "
                "itself via RAG."
            )},
        ]
        pd.DataFrame(dependency_rows).to_excel(writer, index=False, sheet_name="RAG vs Gemini Dependency")
        ws_dep = writer.sheets["RAG vs Gemini Dependency"]
        ws_dep.column_dimensions["A"].width = 42
        ws_dep.column_dimensions["B"].width = 70
        _style_header(ws_dep)

        # ---------------- Sheet 6: Configuration ----------------
        config_rows = [{"Parameter": k, "Value": v} for k, v in run_config_summary.items()]
        config_rows.append({"Parameter": "", "Value": ""})
        config_rows.append({"Parameter": "--- RESULTS ---", "Value": ""})
        config_rows.extend([
            {"Parameter": "Total Documents", "Value": total_docs},
            {"Parameter": "Successful", "Value": success_count},
            {"Parameter": "Errors", "Value": error_count},
            {"Parameter": "Correct Predictions", "Value": correct_count},
            {"Parameter": "Overall Accuracy %", "Value": overall_accuracy},
        ])
        config_rows.append({"Parameter": "", "Value": ""})
        config_rows.append({"Parameter": "--- GEMINI FREE-TIER USAGE (this run) ---", "Value": ""})
        for k, v in gemini_usage_summary.items():
            config_rows.append({"Parameter": k, "Value": v})
        config_rows.append({"Parameter": "", "Value": ""})
        config_rows.append({"Parameter": "--- COST ---", "Value": ""})
        config_rows.append({"Parameter": "Total Cost USD", "Value": 0.0})
        config_rows.append({"Parameter": "Note", "Value": "OCR = free (Tesseract), Gemini = free-tier key. $0 total."})

        if total_docs:
            total_text_pages = int(pd.to_numeric(df.get("he_text_pages_free", 0), errors="coerce").fillna(0).sum())
            total_ocr_pages = int(pd.to_numeric(df.get("he_ocr_pages_paid", 0), errors="coerce").fillna(0).sum())
            config_rows.append({"Parameter": "", "Value": ""})
            config_rows.append({"Parameter": "--- EXTRACTION STATS ---", "Value": ""})
            config_rows.append({"Parameter": "Text Pages (Free - PyMuPDF)", "Value": total_text_pages})
            config_rows.append({"Parameter": "OCR Pages (Free - Tesseract)", "Value": total_ocr_pages})

        pd.DataFrame(config_rows).to_excel(writer, index=False, sheet_name="Configuration")
        ws_config = writer.sheets["Configuration"]
        ws_config.column_dimensions["A"].width = 34
        ws_config.column_dimensions["B"].width = 60
        _style_header(ws_config)

    return str(excel_file)